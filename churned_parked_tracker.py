"""Churned & Parked case tracker -- standalone, re-runnable (upsert only, never
deletes a row). Imports main.py so Usage Streak uses the exact same
_usage_28()/account-resolution code the live dashboard uses -- one source of
truth. Run with: uv run --with "openpyxl>=3.1.5" python churned_parked_tracker.py

Behavior on each run:
  - Every deal CURRENTLY in deal_stage Churned/CS Parked, PLUS every deal ever
    added by a previous run (even if it later moved to a different stage --
    e.g. "brought back" and reactivated), gets its row refreshed from the live
    DB: Stage, CS Owner, Churned Reason, CS Parked Reason, Notes, Usage
    Streak. The row set only grows -- a deal is never removed.
  - "Brought Back" and "Date Brought Back" are pure manual columns: never
    written by this script except to leave them blank for a brand-new row.
  - Any OTHER column a human adds to the sheet is preserved untouched,
    carried forward by Record ID (not by position).
  - Only the "Churned & Parked" sheet is touched; any other sheet in the
    workbook (or a sheet added later) is left exactly as-is.
"""
import sys, os, io, contextlib

_HERE = os.path.dirname(os.path.abspath(__file__))
os.chdir(_HERE)
sys.path.insert(0, _HERE)

_buf = io.StringIO()
with contextlib.redirect_stdout(_buf):
    import main as m
for _line in _buf.getvalue().splitlines():
    print(f"[main] {_line}", file=sys.stderr)

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.cell.text import InlineFont
from openpyxl.cell.rich_text import CellRichText, TextBlock

HS_BASE = "https://app-na2.hubspot.com/contacts/39668252/record/0-3/"
SHEET_NAME = "Churned & Parked"
OUT_PATH = os.path.join(_HERE, "exports", "Churned_Parked_Tracker.xlsx")

# Fixed display columns this script owns (always refreshed from live data,
# except the two manual ones). "Record ID" is the hidden join key.
OWNED_COLS   = ["Deal Name", "Stage", "CS Owner", "Churned Reason",
                "CS Parked Reason", "Notes (aia)", "Usage Streak",
                "Brought Back", "Date Brought Back", "Record ID"]
MANUAL_COLS  = {"Brought Back", "Date Brought Back"}
REFRESH_COLS = [c for c in OWNED_COLS if c not in MANUAL_COLS and c != "Record ID"]

# Usage Streak dot colors -- exact match to grid_server.py's .dot.sync/.on/.seen/.off
# Full 8-digit ARGB (FF alpha prefix) -- a plain 6-digit hex gets padded with 00
# (fully TRANSPARENT) by openpyxl's Color class, which is what made every dot
# render black/invisible in real Excel until this was caught.
GREEN, AMBER, BLUE, GREY = "FF16a34a", "FFeab308", "FF7dd3fc", "FFd8dee6"

def streak_richtext(streak):
    """28 colored dots, left=today .. right=27 days ago -- same order and same
    classification rule as streakHtml() in grid_server.py."""
    blocks = []
    days = (streak or "").split(";")
    for i, day in enumerate(days[:28]):
        p = (day or "").split(",")
        gi = lambda idx: int(p[idx]) if idx < len(p) and p[idx] else 0
        green = gi(2) > 0 or gi(7) > 0
        amber = (gi(1) > 0 or gi(5) > 0 or gi(6) > 0 or gi(8) > 0 or gi(9) > 0 or
                 gi(10) > 0 or gi(11) > 0 or gi(13) > 0 or gi(18) > 0 or gi(19) > 0)
        blue = gi(12) > 0 or gi(4) > 0
        color = GREEN if green else (AMBER if amber else (BLUE if blue else GREY))
        # Trailing space lives INSIDE the same colored run as the dot, not as
        # its own run -- a run whose <t> is pure whitespace needs an explicit
        # xml:space="preserve" that openpyxl's rich-text writer doesn't add,
        # so Excel's strict parser drops it and "repairs" the whole cell.
        sep = " " if i < len(days[:28]) - 1 else ""
        blocks.append(TextBlock(InlineFont(color=color, sz="14"), "●" + sep))
    return CellRichText(blocks) if blocks else ""

def resolve_usage(login_email, poc_email, ev_lu):
    """Usage Streak account resolution: login_email_id first, poc_email as a
    fallback (mirrors the poc_email recovery pattern used elsewhere for
    accounts whose login_email_id doesn't resolve)."""
    acct = m._acct_for(login_email) if login_email else None
    email_used = login_email
    if acct is None and poc_email:
        acct = m._acct_for(poc_email)
        email_used = poc_email
    if acct is None:
        return 0, ""
    active_days, streak, _bq, _bu = m._usage_28(email_used, ev_lu, acct=acct)
    return active_days, streak

def live_row(rec, ev_lu):
    """Build the OWNED-column values for one record_id from m._AIA (current
    live data). Returns None if the record_id no longer exists in aia_live
    (e.g. deleted) -- caller then carries forward the row's last saved state."""
    hit = m._AIA[m._AIA["record_id"] == rec]
    if not len(hit):
        return None
    r = hit.iloc[0]
    _active, streak = resolve_usage(r.get("login_email_id", ""), r.get("poc_email", ""), ev_lu)
    return {
        "Deal Name": r.get("deal_name", "") or "",
        "Stage": r.get("deal_stage", "") or "",
        "CS Owner": r.get("cs_owner", "") or "",
        "Churned Reason": r.get("churned_reason", "") or "",
        "CS Parked Reason": r.get("cs_parked_reason", "") or "",
        "Notes (aia)": r.get("notes", "") or "",
        "Usage Streak": streak,
        "Record ID": rec,
    }

def main():
    ev_lu = m._recent_event_lookup()

    # ── read the existing sheet (if any): header row drives what columns
    # exist (known + any human-added extras), rows keyed by Record ID.
    existing = {}          # record_id -> {col_name: value}
    extra_cols = []        # column names present in the old sheet, unknown to this script
    wb = None
    if os.path.exists(OUT_PATH):
        wb = openpyxl.load_workbook(OUT_PATH)
        if SHEET_NAME in wb.sheetnames:
            ws_old = wb[SHEET_NAME]
            hdr_row = None
            for row in ws_old.iter_rows(min_row=1, max_row=ws_old.max_row):
                vals = [c.value for c in row]
                if "Record ID" in vals:
                    hdr_row = row[0].row
                    headers = vals
                    break
            if hdr_row:
                extra_cols = [h for h in headers if h and h not in OWNED_COLS]
                id_i = headers.index("Record ID")
                for row in ws_old.iter_rows(min_row=hdr_row + 1, max_row=ws_old.max_row):
                    rid = row[id_i].value
                    if rid is None:
                        continue
                    existing[rid] = {headers[i]: row[i].value for i in range(len(headers)) if i < len(row)}
            del wb[SHEET_NAME]   # rebuilt fresh below; other sheets in wb untouched
    if wb is None:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
            del wb["Sheet"]

    # ── union of record_ids: every deal CURRENTLY Churned/CS Parked, plus
    # every record_id ever seen before (upsert, never delete).
    current_ids = set(m._AIA.loc[m._AIA["deal_stage"].isin(["Churned", "CS Parked"]), "record_id"])
    full_ids = current_ids | set(existing.keys())

    out_rows = []
    for rec in full_ids:
        fresh = live_row(rec, ev_lu)
        old = existing.get(rec, {})
        if fresh is not None:
            row = dict(fresh)
            row["Brought Back"] = old.get("Brought Back", "")
            row["Date Brought Back"] = old.get("Date Brought Back", "")
        else:
            # record_id no longer in aia_live (e.g. deleted) -- carry the
            # whole row forward exactly as last saved, nothing to refresh.
            row = {c: old.get(c, "") for c in OWNED_COLS}
        for c in extra_cols:
            row[c] = old.get(c, "")
        out_rows.append(row)

    out_rows.sort(key=lambda r: (r.get("Deal Name") or "").lower())

    # ── write the sheet ──────────────────────────────────────────────────
    ws = wb.create_sheet(SHEET_NAME, 0)
    note = ("Churned & Parked case tracker -- upsert only, rows are never deleted. "
            "Stage/CS Owner/reasons/Notes/Usage Streak refresh from the live DB on "
            "every run; Brought Back + Date Brought Back are yours to fill in and are "
            "never overwritten. Usage Streak: left=today .. right=27 days ago -- "
            "green=accounting sync, amber=other usage, blue=login/view only, grey=none.")
    cols = [c for c in OWNED_COLS if c != "Record ID"] + extra_cols + ["Record ID"]
    ws.append([note])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws["A1"].font = Font(italic=True, size=9, color="555555")
    ws["A1"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[1].height = 45
    ws.append([])
    ws.append(cols)
    header_row = ws.max_row
    HEAD_FILL = PatternFill("solid", fgColor="0F2F52")
    HEAD_FONT = Font(bold=True, color="FFFFFF")
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=header_row, column=c)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    link_font = Font(color="1A7FC4", underline="single")
    # the GREY dot color (#d8dee6) is nearly invisible on Excel's pure-white
    # default cell fill -- a faint background gives it contrast without
    # changing the dot colors themselves (which must match the dashboard).
    STREAK_FILL = PatternFill("solid", fgColor="FFF8FAFC")
    name_col = cols.index("Deal Name") + 1
    streak_col = cols.index("Usage Streak") + 1
    for r in out_rows:
        vals = [r.get(c, "") for c in cols]
        ws.append(vals)
        rix = ws.max_row
        rid = r.get("Record ID")
        if rid and r.get("Deal Name"):
            cell = ws.cell(row=rix, column=name_col)
            cell.hyperlink = HS_BASE + str(rid)
            cell.font = link_font
        streak_cell = ws.cell(row=rix, column=streak_col)
        streak_cell.fill = STREAK_FILL
        streak_val = r.get("Usage Streak")
        if streak_val:
            streak_cell.value = streak_richtext(streak_val)

    widths = {"Deal Name": 40, "Stage": 20, "CS Owner": 18, "Churned Reason": 28,
              "CS Parked Reason": 28, "Notes (aia)": 34, "Brought Back": 14,
              "Date Brought Back": 16, "Usage Streak": 65}
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 20)
    rid_col = cols.index("Record ID") + 1
    ws.column_dimensions[get_column_letter(rid_col)].hidden = True
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(cols))}{ws.max_row}"

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")
    print(f"Rows: {len(out_rows)} (currently Churned/CS Parked: {len(current_ids)}, "
          f"carried forward from before: {len(full_ids - current_ids)})")
    if extra_cols:
        print(f"Preserved extra columns: {extra_cols}")

if __name__ == "__main__":
    main()
